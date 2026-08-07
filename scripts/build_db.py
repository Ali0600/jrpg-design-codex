from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

wb = Workbook()

HDR_FILL = PatternFill('solid', start_color='1F3A5F')
HDR_FONT = Font(name='Arial', bold=True, color='FFFFFF', size=11)
BODY_FONT = Font(name='Arial', size=10)
TITLE_FONT = Font(name='Arial', bold=True, size=14, color='1F3A5F')
SUB_FONT = Font(name='Arial', italic=True, size=10, color='555555')
YES_FILL = PatternFill('solid', start_color='D9EAD3')
WRAP = Alignment(wrap_text=True, vertical='top')
THIN = Border(*[Side(style='thin', color='CCCCCC')]*4)

def style_header(ws, row, ncols):
    for c in range(1, ncols+1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HDR_FILL; cell.font = HDR_FONT
        cell.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')

def fill_rows(ws, start_row, rows, ncols):
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            cell = ws.cell(row=start_row+i, column=j+1, value=v)
            cell.font = BODY_FONT; cell.alignment = WRAP; cell.border = THIN

# ---------------- README ----------------
ws = wb.active; ws.title = 'README'
ws['A1'] = 'JRPG Design Research Database'; ws['A1'].font = TITLE_FONT
ws['A2'] = 'Your core design pillar: A world that rewards exploration — everything has upgrades.'
ws['A2'].font = SUB_FONT
readme = [
    ('', ''),
    ('HOW TO USE THIS WORKBOOK', ''),
    ('1. Games sheet', 'Your research roster. Add every game you want to study. Track status and why it matters to your design.'),
    ('2. Mechanics sheet', 'The core database. One row per mechanic. When you play/research a game, log each mechanic you find, rate it 1-5, decide Yes/Maybe/No on adopting it, and write how YOU would adapt it.'),
    ('3. Design Pillars sheet', 'Your north star. Every mechanic should be checked against these. Edit freely as your vision sharpens.'),
    ('4. Summary sheet', 'Auto-counts what you have logged and what you have committed to building.'),
    ('', ''),
    ('WORKFLOW FOR EACH NEW GAME', ''),
    ('Step 1', 'Add the game to the Games sheet with Status = To Research.'),
    ('Step 2', 'Play it (or watch/read deep dives) with the Mechanics sheet open. Log mechanics as you notice them.'),
    ('Step 3', 'For every mechanic, always fill "Reward Loop" — what does the player DO, what do they GET, and why does it feel good? This is the column that matters most for your exploration-reward pillar.'),
    ('Step 4', 'Rate it and mark "Want in My Game?". For Yes/Maybe, write Adaptation Notes: how the idea translates to YOUR game, stripped of the source game specifics.'),
    ('Step 5', 'Set the game Status to Researched. Periodically re-read all "Yes" rows together — that cluster IS your game design document seed.'),
    ('', ''),
    ('CATEGORY DEFINITIONS', ''),
    ('Exploration & Rewards', 'Anything about discovering the world and being paid for it (your pillar — e.g., Elden Ring paintings, FF7R lifesprings).'),
    ('Progression & Upgrades', 'Character/gear growth systems (e.g., Materia, weapon levels, folios).'),
    ('Combat', 'Battle systems and moment-to-moment fight mechanics.'),
    ('Economy & Currency', 'Currencies, shops, crafting, and what they gate.'),
    ('Traversal', 'How the player moves; mounts, abilities that open up the map.'),
    ('Side Content & Minigames', 'Optional activities and their reward hooks.'),
    ('Social & Relationships', 'Bond/affinity systems between characters.'),
    ('Meta & Replayability', 'NG+, hard modes, chapter select, completion systems.'),
]
r = 4
for a, b in readme:
    ws.cell(row=r, column=1, value=a).font = Font(name='Arial', bold=(b=='' and a!=''), size=10)
    c2 = ws.cell(row=r, column=2, value=b); c2.font = BODY_FONT; c2.alignment = WRAP
    r += 1
ws.column_dimensions['A'].width = 26
ws.column_dimensions['B'].width = 110

# ---------------- Games ----------------
ws = wb.create_sheet('Games')
gcols = ['Game', 'Year', 'Developer', 'Genre / Style', 'Status', 'Priority', 'Why Research It (what it can teach you)']
ws.append(gcols); style_header(ws, 1, len(gcols))
games = [
 ['Final Fantasy VII Rebirth', 2024, 'Square Enix', 'Action JRPG, semi-open world', 'Researched', 'High',
  'Layered upgrade systems (Materia, weapon levels, folios) + World Intel exploration-reward structure. Your reference model for "everything has upgrades".'],
 ['Elden Ring', 2022, 'FromSoftware', 'Action RPG, open world', 'Researching', 'High',
  'Gold standard for curiosity-driven exploration: paintings, hidden walls, every landmark pays off. No checklist — the world itself is the quest log.'],
 ['Final Fantasy X', 2001, 'Square', 'Turn-based JRPG', 'To Research', 'Medium',
  'Sphere Grid: iconic node-based progression that makes every battle feed a visible upgrade path.'],
 ['Persona 5 Royal', 2019, 'Atlus', 'Turn-based JRPG, social sim', 'To Research', 'Medium',
  'Confidant system: relationships ARE upgrades. Calendar structure makes all activities feel rewarding.'],
 ['Xenoblade Chronicles 3', 2022, 'Monolith Soft', 'Action JRPG, open zones', 'To Research', 'Medium',
  'Massive explorable zones where discovering landmarks/areas grants EXP directly — exploring literally levels you up.'],
 ['Dragon Quest XI', 2017, 'Square Enix', 'Turn-based JRPG', 'To Research', 'Low',
  'Classic JRPG reward pacing; forge/crafting minigame; recipe books found by exploring.'],
 ['The Legend of Zelda: Tears of the Kingdom', 2023, 'Nintendo', 'Action-adventure, open world', 'To Research', 'Medium',
  'Shrines/Korok seeds: micro-rewards that make every corner of the map worth checking; rewards convert into upgrades (hearts/stamina/inventory).'],
 ['Chained Echoes', 2022, 'Matthias Linda', 'Indie 16-bit style JRPG', 'To Research', 'Low',
  'Solo-dev JRPG (relevant to your scope!). Reward Board system: exploration and quest achievements directly unlock gear/upgrades.'],
 ['Sea of Stars', 2023, 'Sabotage Studio', 'Indie turn-based JRPG', 'To Research', 'Low',
  'Small-team JRPG benchmark: relic system, cooking, and exploration secrets. Good scope reference for an indie JRPG.'],
]
fill_rows(ws, 2, games, len(gcols))
widths = [30, 8, 18, 28, 14, 10, 90]
for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = 'A2'
ws.auto_filter.ref = f'A1:G{1+len(games)}'
dv_status = DataValidation(type='list', formula1='"To Research,Researching,Researched"', allow_blank=True)
dv_prio = DataValidation(type='list', formula1='"High,Medium,Low"', allow_blank=True)
ws.add_data_validation(dv_status); ws.add_data_validation(dv_prio)
dv_status.add('E2:E200'); dv_prio.add('F2:F200')

# ---------------- Mechanics ----------------
ws = wb.create_sheet('Mechanics')
mcols = ['ID', 'Game', 'Mechanic', 'Category', 'How It Works', 'Reward Loop (do X -> get Y -> feels good because Z)',
         'Rating (1-5)', 'Want in My Game?', 'Adaptation Notes (how it fits YOUR game)']
ws.append(mcols); style_header(ws, 1, len(mcols))

FF = 'Final Fantasy VII Rebirth'
ER = 'Elden Ring'
mechs = [
 # --- FF7R: Progression & Upgrades ---
 [FF, 'Materia System', 'Progression & Upgrades',
  'Orbs slotted into weapons/armor grant spells, abilities, and passives. Materia earns AP from battles and levels up (stronger spells, new effects). Linked slots let two materia combine (e.g., Fire + Elemental = fire-infused sword).',
  'Fight -> materia gains AP -> spell tiers unlock -> your loadout is a build you authored. Slots are scarce, so every choice is meaningful; linked slots reward experimentation.',
  '', 'Yes', 'You flagged this as a favorite. Core idea to steal: upgrades are SOCKETED and TRANSFERABLE, not welded to one item — the player owns a growing library of upgrades and expresses themselves through slotting.'],
 [FF, 'Weapon Levels + Weapon Skills', 'Progression & Upgrades',
  'Every weapon a character owns levels up together as the character earns SP (total SP earned = weapon level). Each level auto-grants stat boosts, new equippable Weapon Skill perks, and sometimes extra materia slots. Skill slots are limited, so you pick a subset.',
  'Earn SP anywhere -> ALL weapons improve -> old weapons never become trash; new skill slots create fresh decisions each level.',
  '', 'Yes', 'You flagged this as a favorite. Key insight: no dead weapons. Every weapon stays viable and keeps gaining perks, so finding a new one is about its unique ability, not raw numbers.'],
 [FF, 'Weapon Abilities + Proficiency', 'Progression & Upgrades',
  'Each weapon has a signature ability. Using it and completing its proficiency objective teaches the ability PERMANENTLY, even after unequipping.',
  'Find weapon -> use its move enough -> keep the move forever -> every weapon pickup is a guaranteed permanent upgrade, not just a stat stick.',
  '', 'Maybe', 'Very compatible with your pillar: exploration finds weapons, weapons teach permanent moves. Consider making proficiency objectives mini skill-challenges.'],
 [FF, 'Folios (Skill Trees)', 'Progression & Upgrades',
  'Node-based skill boards per character. Spend SP on nodes for stat boosts, new abilities, synergy moves, and Lv3 limit breaks. Node waves unlock via Party Level; some nodes gated by Weapon Level. Respec is free. Must visit bookstores/automats in the world to edit.',
  'Every activity feeds SP/Party Level -> visible new nodes light up -> constant "one more unlock" pull. Free respec removes fear of mistakes.',
  '', '', 'Note the friction choice: editing folios requires visiting a physical location in the world — even menu actions pull you into exploring towns.'],
 [FF, 'SP as a triple-duty resource', 'Progression & Upgrades',
  'One resource (SP) drives three systems: total earned sets Weapon Level, spendable pool buys Folio nodes, and combined party SP sets Party Level.',
  'Any SP source (level up, manuscripts, quests) visibly moves THREE progress bars at once -> everything you do feels multiplicatively rewarding.',
  '', 'Maybe', 'Elegant trick for "everything has upgrades": make one currency ripple through multiple systems so every reward hits 2-3 progress meters.'],
 [FF, 'Manuscripts (SP collectibles)', 'Progression & Upgrades',
  'Collectible books that grant SP to a specific character. Found in side quests, minigame rewards, moogle shops, and hard-mode boss replays.',
  'Do optional content -> concrete permanent character power -> side content never feels skippable-filler.',
  '', '', 'Model for rewarding side content with PROGRESSION currency instead of consumables.'],
 [FF, 'Party Level / Party EXP', 'Progression & Upgrades',
  'A separate account-wide level raised almost exclusively by side content (World Intel, quests). Unlocks waves of folio nodes for EVERY character.',
  'Explore -> whole party grows (not just active members) -> exploration is never wasted on benched characters.',
  '', '', 'Solves the classic JRPG problem of benched party members falling behind.'],
 # --- FF7R: Combat ---
 [FF, 'ATB Hybrid Combat', 'Combat',
  'Real-time action (attack/dodge/block) charges ATB gauges; open the command menu to slow time ("Tactical Mode") and spend ATB on abilities, magic, and items.',
  'Aggression is the engine: attacking earns the resource that unlocks your cool moves -> constant push-forward incentive with strategic pauses.',
  '', '', 'The core loop: normal actions charge a resource, the resource buys spectacular actions.'],
 [FF, 'Pressure & Stagger', 'Combat',
  'Exploiting weaknesses or specific tactics "pressures" enemies; pressured enemies fill a stagger gauge faster; staggered enemies take massively boosted damage. Assess materia reveals each enemy\'s pressure conditions. Bosses have destructible parts with own health bars.',
  'Learn the enemy -> execute its puzzle -> huge damage window payoff -> every fight is a knowledge check, not a stat check.',
  '', '', 'Turns combat itself into exploration (of enemy design). Assess = in-combat "scanning" reward.'],
 [FF, 'Synergy Skills & Abilities', 'Combat',
  'Free two-character team-up skills (no ATB cost) that build ATB for both users. Spending ATB builds synergy charges; full charges unlock cinematic paired Synergy Abilities with bonus effects.',
  'Cooperation compounds: small team moves feed the resource -> big team moves -> party feels like a band, not solo heroes rotating.',
  '', '', 'Also gated by the relationship/folio systems — combat power expresses character bonds.'],
 [FF, 'Switchable 3-character party', 'Combat',
  'Direct control of one of three frontliners, hot-swappable anytime; each character is a genuinely different playstyle (Cloud stance-dance, Barret ranged, Tifa combo-rush, Aerith zoning wards).',
  'Mastery depth: one combat system, seven different ways to play it -> replay value inside every battle.',
  '', '', ''],
 # --- FF7R: Exploration & Rewards (the pillar) ---
 [FF, 'World Intel (umbrella checklist)', 'Exploration & Rewards',
  'Each open region has categories of discoverable activities (towers, lifesprings, fiend hunts, shrines, digs, moogles, protorelics). ALL of them pay out the same two meta-rewards — regional Data Points and Party EXP — plus activity-specific rewards.',
  'Any activity -> guaranteed dual payout (currency + party growth) + its own unique reward -> no exploration is ever "empty calories".',
  '', '', 'The architecture to study: many activity TYPES, one unified reward SPINE. Consistency makes the player trust that exploring always pays.'],
 [FF, 'Remnawave Towers', 'Exploration & Rewards',
  'Climbable towers (guarded by enemies) that, when activated, reveal nearby points of interest on the map and become fast-travel points.',
  'Small effort -> map knowledge + convenience -> reduces friction for all future exploration in the region.',
  '', '', 'Classic Ubisoft-style tower. Decide: do you want revealed icons (guided) or Elden Ring-style unguided discovery? See ER rows.'],
 [FF, 'Lifesprings + Springseeker Owls', 'Exploration & Rewards',
  'Glowing owls in the wild lead you to hidden mako springs. Activating one grants regional lore, a materials cache, and progressively unlocks Excavation sites and the region\'s secret boss (Classified Intel).',
  'Follow a curious creature -> discover a secret spot -> get lore + loot + NEW secrets unlocked -> discovery breeds more discovery.',
  '', 'Maybe', 'Great "organic guide" pattern: instead of map markers, a living creature invites you off the road. Rewards that unlock further rewards = chained discovery.'],
 [FF, 'Fiend Intel (elite hunts w/ bonus objectives)', 'Exploration & Rewards',
  'Rare elite monster variants in the overworld. Each fight has 3 optional objectives (e.g., stagger it, win before X move) that multiply the rewards and unlock combat simulator content.',
  'Optional challenge tiers inside the encounter -> skilled play is directly paid -> same content rewards casual and expert players differently.',
  '', '', ''],
 [FF, 'Classified Intel (regional secret boss)', 'Exploration & Rewards',
  'Completing all lifesprings in a region reveals a hidden legendary monster. It drops materials found nowhere else and unlocks new enemy-skill learning fights.',
  'Thorough exploration -> exclusive climactic challenge -> exclusive rewards -> completionism gets a boss fight as its trophy.',
  '', 'Maybe', 'Strong pillar fit: the REWARD for exploring everything is more GAME (a secret boss), not just an item.'],
 [FF, 'Summon Sanctuaries (Divine Intel)', 'Exploration & Rewards',
  'Each region hides 3 shrines tied to its summon deity. Solving each shrine\'s memory minigame weakens the summon\'s trial boss AND levels up the summon materia (finding all 3 maxes it at 4 stars with a new ability).',
  'Explore shrines -> the upcoming boss gets easier AND the prize gets stronger -> exploration converts directly into both accessibility and power.',
  '', 'Yes', 'One of the best "world exploration = upgrades" designs in the game. The same collectible tunes difficulty down and power up simultaneously.'],
 [FF, 'Excavation Intel (chocobo digging)', 'Exploration & Rewards',
  'Lifesprings reveal dig sites; ride a chocobo that sniffs out buried treasure, including Transmuter Chips that permanently expand your crafting recipe list.',
  'Discovery -> buried treasure -> permanent crafting expansion -> loot that upgrades a SYSTEM, not just your inventory.',
  '', 'Maybe', 'Steal this: some exploration rewards should upgrade your capabilities/systems (recipes, tools) rather than hand over items.'],
 [FF, 'Protorelic questlines', 'Exploration & Rewards',
  'Each region has a 4-stage themed quest chain (different genre of activity per region) hunting a mysterious relic. Collecting all protorelics across the game unlocks a powerful secret summon.',
  'Regional stories chain into a game-spanning meta-collection -> long-horizon reward for consistent explorers.',
  '', '', 'Pattern: local collectibles that also feed ONE big global prize.'],
 [FF, 'Mogstools / Moogle Emporium', 'Exploration & Rewards',
  'Hidden moogle houses with a herding minigame. Each clear raises your Emporium level, unlocking new stock (manuscripts, materia, gear) purchasable with Moogle Medals earned from minigames.',
  'Find secret spot -> play toy challenge -> a SHOP grows -> the reward is expanded purchasing options, which chain into further upgrades.',
  '', '', ''],
 [FF, 'Chocobo Intel (regional mounts w/ unique traversal)', 'Traversal',
  'Each region has a stealth challenge to capture its native chocobo, which has a region-specific traversal power (gliding, wall-climbing, water-walking) required to reach that region\'s secrets.',
  'Earn the mount -> previously visible-but-unreachable places open up -> the map itself is an upgrade tree.',
  '', 'Yes', 'Metroidvania logic in an open world: traversal abilities ARE upgrades, and terrain is the lock. Perfect fit for exploration-first design.'],
 # --- FF7R: Economy ---
 [FF, 'Region-locked Data Points -> Materia development', 'Economy & Currency',
  'All World Intel pays Data Points, but they are locked to the region they were earned in. Spent with Chadley to develop unique materia per region.',
  'Forces genuine engagement with EACH region (can\'t grind one zone for everything) -> every new area has its own exclusive upgrade shelf.',
  '', 'Maybe', 'Interesting anti-grind lever: regional currencies keep every area relevant. Could feel restrictive — decide consciously.'],
 [FF, 'Item Transmuter (portable crafting)', 'Economy & Currency',
  'Craft consumables/gear anywhere from gathered materials. Crafting new item types grants transmuter EXP, leveling it up to unlock more recipes; chips found by exploring expand it further.',
  'Gathering while exploring -> crafting -> the crafting SYSTEM levels up too -> even your tools have an upgrade track.',
  '', 'Yes', 'Literal embodiment of "everything has upgrades" — the crafting menu itself has XP and levels.'],
 # --- FF7R: Side content ---
 [FF, 'Queen\'s Blood (collectible card game)', 'Side Content & Minigames',
  'Full deck-builder card minigame. Cards are acquired from shops, quest rewards, and beating opponents scattered across the whole world; a dedicated storyline runs through it.',
  'Every town hides new opponents/cards -> exploration feeds a parallel collection/mastery loop independent of combat power.',
  '', '', 'A minigame with its own progression + collection + narrative = a second game keeping the world sticky.'],
 [FF, 'Gold Saucer & minigame economy (GP)', 'Side Content & Minigames',
  'Theme-park hub of minigames paying GP, a dedicated currency for exclusive prizes; dozens of one-off minigames also gate manuscripts and gear throughout the game.',
  'Play variety content -> exclusive currency -> exclusive prizes -> changes of pace still feed the reward spine.',
  '', '', ''],
 [FF, 'Bond / relationship system', 'Social & Relationships',
  'Dialogue choices, side quests, and synergy usage raise hidden affinity per companion; affinity affects scenes (notably the Gold Saucer date) and is boosted by doing character-linked side content.',
  'Optional content deepens relationships -> emotional payoff scenes -> rewards can be narrative, not just numeric.',
  '', '', 'Reminder for your design: "reward" can be story, intimacy, and lore — not only stats.'],
 [FF, 'Chapter Select / Hard Mode NG+', 'Meta & Replayability',
  'Finishing the game unlocks chapter select with all progress retained. Hard mode bans items, gives bosses new moves, and drops exclusive manuscripts needed to max folios.',
  'Full character completion REQUIRES mastering hard mode -> replay is an upgrade path, not a reset.',
  '', '', ''],
 # --- Elden Ring ---
 [ER, 'Paintings -> vista treasure', 'Exploration & Rewards',
  'Find a painting in a ruin; it depicts a real location somewhere in the world. Travel there, identify the painted vantage point, and a ghost appears with a reward on the spot.',
  'Discovery creates a riddle -> the WORLD is the answer key -> solving it with your own eyes pays off with an item. Reward for paying attention to scenery.',
  '', 'Yes', 'You flagged this as a favorite. The genius: it rewards knowing the world, not clearing a marker. Your version could use sketches, photos, songs, or NPC rumors that describe places visually.'],
 [ER, 'Unmarked world density (every landmark pays off)', 'Exploration & Rewards',
  'No quest markers or checklists. Every ruin, cave, and odd silhouette on the horizon reliably contains SOMETHING (boss, talisman, crafting book, NPC). Trust is built through consistency.',
  'See interesting shape -> go there -> always rewarded -> curiosity itself becomes the quest log.',
  '', 'Maybe', 'The opposite philosophy to FF7R\'s World Intel checklist. Decide where your game sits on the guided <-> unguided spectrum; you can hybridize.'],
 [ER, 'Golden Seeds & Sacred Tears (exploration -> healing upgrades)', 'Exploration & Rewards',
  'Glowing saplings and church statues scattered across the map hold seeds/tears that directly upgrade your healing flask count and potency.',
  'Spot a landmark -> walk over -> permanent survivability upgrade -> the most vital upgrade resource is placed as pure exploration bait.',
  '', 'Maybe', 'Pattern: put your game\'s most universally-wanted upgrade out in the world as visible landmarks, not in shops.'],
 [ER, 'Illusory walls & secret-behind-secret', 'Exploration & Rewards',
  'Some walls are fake and hide rooms, sometimes chaining multiple layers deep. Message system lets players hint to each other.',
  'Paranoid curiosity -> occasional jackpot -> stories players tell each other. Rewards the "what if I check?" instinct.',
  '', '', ''],
]
for i, m in enumerate(mechs, 1):
    row = [f'M{i:03d}'] + m
    for j, v in enumerate(row):
        cell = ws.cell(row=1+i, column=j+1, value=v)
        cell.font = BODY_FONT; cell.alignment = WRAP; cell.border = THIN
    if m[6] == 'Yes':
        ws.cell(row=1+i, column=8).fill = YES_FILL

mwidths = [7, 24, 30, 22, 60, 60, 10, 12, 60]
for i, w in enumerate(mwidths, 1): ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = 'A2'
last = 1 + len(mechs)
ws.auto_filter.ref = f'A1:I{last}'
dv_cat = DataValidation(type='list', formula1='"Exploration & Rewards,Progression & Upgrades,Combat,Economy & Currency,Traversal,Side Content & Minigames,Social & Relationships,Meta & Replayability"', allow_blank=True)
dv_rate = DataValidation(type='list', formula1='"1,2,3,4,5"', allow_blank=True)
dv_want = DataValidation(type='list', formula1='"Yes,Maybe,No"', allow_blank=True)
for dv in (dv_cat, dv_rate, dv_want): ws.add_data_validation(dv)
dv_cat.add('D2:D1000'); dv_rate.add('G2:G1000'); dv_want.add('H2:H1000')

# ---------------- Design Pillars ----------------
ws = wb.create_sheet('Design Pillars')
pcols = ['#', 'Pillar', 'What it means', 'Test question for any mechanic']
ws.append(pcols); style_header(ws, 1, len(pcols))
pillars = [
 [1, 'The world rewards exploration', 'Every place worth looking at contains something worth finding. No empty vistas, no dead ends without a payoff.',
  'If a player wanders off the path here, do they get something? (item, lore, upgrade, secret, or more game)'],
 [2, 'Everything has upgrades', 'Not just characters: weapons, tools, shops, crafting, mounts, even minigames should have their own growth track.',
  'Does this system have its own progress bar? Can the player make it better over time?'],
 [3, 'Discovery breeds discovery', 'Finding one secret should unlock, hint at, or reveal the next one (FF7R lifesprings -> dig sites -> secret boss).',
  'Does this reward chain into another reward?'],
 [4, 'Rewards are layered', 'One action should pay out on multiple meters at once (FF7R: any intel = Data Points + Party EXP + its own prize).',
  'How many progress systems does completing this touch? (aim for 2+)'],
 [5, 'Knowledge is a reward', 'Lore, enemy weaknesses, map understanding, and riddle answers (Elden Ring paintings) are treasures too.',
  'Can a player be rewarded here just for paying attention?'],
]
fill_rows(ws, 2, pillars, len(pcols))
for i, w in enumerate([5, 30, 70, 60], 1): ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = 'A2'

# ---------------- Summary ----------------
ws = wb.create_sheet('Summary')
ws['A1'] = 'Research Summary'; ws['A1'].font = TITLE_FONT
rows = [
 ('Games in roster', "=COUNTA(Games!A2:A200)"),
 ('Games researched', '=COUNTIF(Games!E2:E200,"Researched")'),
 ('Games still to research', '=COUNTIF(Games!E2:E200,"To Research")'),
 ('Mechanics logged', "=COUNTA(Mechanics!A2:A1000)"),
 ('Mechanics you want in your game (Yes)', '=COUNTIF(Mechanics!H2:H1000,"Yes")'),
 ('Mechanics under consideration (Maybe)', '=COUNTIF(Mechanics!H2:H1000,"Maybe")'),
 ('Exploration & Rewards mechanics logged', '=COUNTIF(Mechanics!D2:D1000,"Exploration & Rewards")'),
 ('Progression & Upgrades mechanics logged', '=COUNTIF(Mechanics!D2:D1000,"Progression & Upgrades")'),
]
r = 3
for label, f in rows:
    a = ws.cell(row=r, column=1, value=label); a.font = BODY_FONT
    b = ws.cell(row=r, column=2, value=f); b.font = Font(name='Arial', bold=True, size=11)
    r += 1
ws.column_dimensions['A'].width = 45
ws.column_dimensions['B'].width = 12

wb.save('/home/claude/gamedb/JRPG_Design_Research_Database.xlsx')
print('saved')
